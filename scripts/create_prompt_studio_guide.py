from __future__ import annotations

import os
import textwrap
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Prompt Studio 团队使用说明.docx"
ASSET_DIR = Path(os.environ.get("TEMP", str(ROOT / "tmp"))) / "prompt-studio-guide-assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT_REGULAR = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"

TEAL = "0B6B61"
TEAL_DARK = "08574F"
TEAL_SOFT = "E9F5F2"
CORAL = "E46F51"
CORAL_SOFT = "FFF0EB"
GOLD = "F5C76B"
INK = "172522"
MUTED = "62706C"
BORDER = "D9E1DE"
LIGHT = "F4F6F5"
LIGHTER = "F7F9F8"
BLUE = "3D79B8"
BLUE_SOFT = "ECF3FA"
GREEN = "18805D"
GREEN_SOFT = "EAF7F1"
AMBER = "A8620A"
AMBER_SOFT = "FFF4DF"
RED = "B83F52"
RED_SOFT = "FCECEF"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def draw_logo(path: Path) -> None:
    canvas = Image.new("RGBA", (256, 256), (255, 255, 255, 0))
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((8, 8, 248, 248), radius=58, fill="#0B6B61")
    draw.line((75, 76, 126, 128, 75, 180), fill="#FFFFFF", width=24, joint="curve")
    draw.line((154, 180, 202, 180), fill="#F28562", width=24)
    draw.ellipse((188, 58, 218, 88), fill="#F5C76B")
    canvas.save(path)


def centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font: ImageFont.FreeTypeFont, fill: str) -> None:
    left, top, right, bottom = box
    lines = textwrap.wrap(text, width=max(7, int((right - left) / max(font.size * 0.95, 1))))
    line_height = font.size + 8
    total_height = len(lines) * line_height - 8
    cursor_y = top + ((bottom - top) - total_height) / 2
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=font)
        cursor_x = left + ((right - left) - (bounds[2] - bounds[0])) / 2
        draw.text((cursor_x, cursor_y), line, font=font, fill=fill)
        cursor_y += line_height


def draw_flow(path: Path) -> None:
    width, height = 1800, 620
    image = Image.new("RGB", (width, height), "#F4F6F5")
    draw = ImageDraw.Draw(image)
    draw.text((72, 42), "Prompt Studio 的标准工作流", font=pil_font(34, True), fill="#172522")
    draw.text((74, 93), "从可复用资源到可复核证据，所有关键条件都可以被团队共同维护", font=pil_font(20), fill="#62706C")

    cards = [
        ("01", "准备资源", "角色 · 世界观\n测试集 · 模型"),
        ("02", "Prompt 调试", "左右窗口\n实时对话"),
        ("03", "保存版本", "新建不可变版本\n发布或切回"),
        ("04", "运行评测", "旧版 / 新版\n固定条件对比"),
        ("05", "人工复核", "逐条证据\n团队打分"),
        ("06", "发布结论", "查看差异\n沉淀下一版"),
    ]
    card_width, card_height, gap = 248, 270, 34
    start_x, start_y = 72, 190
    for index, (number, title, detail) in enumerate(cards):
        x = start_x + index * (card_width + gap)
        draw.rounded_rectangle((x, start_y, x + card_width, start_y + card_height), radius=22, fill="#FFFFFF", outline="#D9E1DE", width=3)
        accent = CORAL if index == 1 else TEAL if index in (0, 2, 3, 5) else GOLD
        draw.rounded_rectangle((x + 24, start_y + 24, x + 76, start_y + 76), radius=16, fill=f"#{accent}")
        draw.text((x + 38, start_y + 34), number, font=pil_font(22, True), fill="#FFFFFF" if accent != GOLD else "#6E500A")
        draw.text((x + 24, start_y + 108), title, font=pil_font(25, True), fill="#172522")
        centered_text(draw, (x + 20, start_y + 154, x + card_width - 20, start_y + 244), detail, pil_font(18), "#62706C")
        if index < len(cards) - 1:
            arrow_start = x + card_width + 8
            arrow_end = x + card_width + gap - 8
            arrow_y = start_y + card_height / 2
            draw.line((arrow_start, arrow_y, arrow_end, arrow_y), fill="#8BBDB4", width=5)
            draw.polygon([(arrow_end, arrow_y), (arrow_end - 13, arrow_y - 10), (arrow_end - 13, arrow_y + 10)], fill="#8BBDB4")
    image.save(path, quality=95)


def draw_roles(path: Path) -> None:
    width, height = 1800, 900
    image = Image.new("RGB", (width, height), "#F4F6F5")
    draw = ImageDraw.Draw(image)
    draw.text((72, 38), "不同角色的使用路径", font=pil_font(34, True), fill="#172522")
    draw.text((74, 90), "按职责进入不同模块，权限由服务端实时校验", font=pil_font(20), fill="#62706C")

    cards = [
        ("负责人", "管理团队与项目", "创建账号\n新建项目\n分配项目角色\n发布版本", TEAL),
        ("编辑者", "建设与运行", "维护 Prompt 和资源\n创建新版本\n配置测试集\n运行评测", CORAL),
        ("评审者", "逐条复核证据", "查看运行结果\n对比两侧回复\n逐案例打分\n提交结论和意见", BLUE),
        ("观察者", "只读了解进展", "查看项目\n阅读 Prompt 版本\n查看评测证据\n参与多人聊天体验", AMBER),
    ]
    card_width, card_height, gap = 386, 565, 28
    start_x, start_y = 72, 180
    for index, (role, subtitle, steps, accent) in enumerate(cards):
        x = start_x + index * (card_width + gap)
        draw.rounded_rectangle((x, start_y, x + card_width, start_y + card_height), radius=22, fill="#FFFFFF", outline="#D9E1DE", width=3)
        draw.rounded_rectangle((x, start_y, x + card_width, start_y + 116), radius=22, fill=f"#{accent}")
        draw.rectangle((x, start_y + 88, x + card_width, start_y + 116), fill=f"#{accent}")
        draw.text((x + 26, start_y + 27), role, font=pil_font(30, True), fill="#FFFFFF" if accent != GOLD else "#6E500A")
        draw.text((x + 27, start_y + 72), subtitle, font=pil_font(18), fill="#FFFFFF" if accent != GOLD else "#6E500A")
        draw.text((x + 28, start_y + 153), "推荐路径", font=pil_font(18, True), fill="#172522")
        step_lines = steps.split("\n")
        current_y = start_y + 206
        for step_index, step in enumerate(step_lines, 1):
            draw.ellipse((x + 30, current_y + 2, x + 62, current_y + 34), fill=f"#{accent}")
            number_color = "#6E500A" if accent == GOLD else "#FFFFFF"
            draw.text((x + 41, current_y + 5), str(step_index), font=pil_font(16, True), fill=number_color)
            draw.text((x + 82, current_y + 3), step, font=pil_font(20), fill="#40514C")
            if step_index < len(step_lines):
                draw.line((x + 46, current_y + 38, x + 46, current_y + 73), fill="#D9E1DE", width=3)
            current_y += 82
        draw.rounded_rectangle((x + 26, start_y + card_height - 78, x + card_width - 26, start_y + card_height - 28), radius=12, fill="#F7F9F8")
        note = "可管理" if role == "负责人" else "按需协作"
        centered_text(draw, (x + 28, start_y + card_height - 75, x + card_width - 28, start_y + card_height - 31), note, pil_font(17, True), f"#{accent}")
    image.save(path, quality=95)


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), str(indent))
    table_indent.set(qn("w:type"), "dxa")
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(widths[index]))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_paragraph_border(paragraph, color: str = TEAL, size: str = "14", space: str = "6") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def style_callout_paragraph(paragraph, fill: str, accent: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    set_run_font(run, size=9, color=MUTED)


def set_keep(paragraph, keep_with_next: bool = False, keep_together: bool = False) -> None:
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = keep_together


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=10.5, color=INK, bold=True)
        body = paragraph.add_run(text[len(bold_lead):])
        set_run_font(body, size=10.5, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 11.5}[level], color=TEAL if level < 3 else TEAL_DARK, bold=True)


def add_callout(doc: Document, label: str, text: str, fill: str = TEAL_SOFT, accent: str = TEAL) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    style_callout_paragraph(paragraph, fill, accent)
    label_run = paragraph.add_run(label + "  ")
    set_run_font(label_run, size=10, color=accent, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10, color=INK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], accent: str = TEAL) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header_row = table.rows[0]
    repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, accent)
        set_cell_border(cell, color=accent, size="6")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(header)
        set_run_font(run, size=9.5, color=WHITE if accent != GOLD else "6E500A", bold=True)
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            set_cell_border(cell)
            set_cell_shading(cell, WHITE if len(table.rows) % 2 else LIGHTER)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, color=INK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_image(doc: Document, path: Path, width: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", caption)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(10)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=8.5, color=MUTED, italic=True)


def add_link_like(paragraph, text: str, color: str = TEAL, bold: bool = True) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=color, bold=bold)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after in ((1, 16, 18, 10), (2, 13, 14, 7), (3, 11.5, 10, 5)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(TEAL if level < 3 else TEAL_DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Kicker" not in [style.name for style in document.styles]:
        kicker = document.styles.add_style("Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = document.styles["Kicker"]
    kicker.font.name = "Segoe UI"
    kicker._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    kicker.font.size = Pt(9)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(CORAL)
    kicker.paragraph_format.space_after = Pt(8)


def build_document() -> None:
    logo_path = ASSET_DIR / "prompt-studio-logo.png"
    flow_path = ASSET_DIR / "prompt-studio-flow.png"
    roles_path = ASSET_DIR / "prompt-studio-roles.png"
    draw_logo(logo_path)
    draw_flow(flow_path)
    draw_roles(roles_path)

    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    document.core_properties.title = "Prompt Studio 团队使用说明"
    document.core_properties.subject = "Prompt 评测、调试、版本管理与团队协作"
    document.core_properties.author = "Prompt Studio"
    document.core_properties.keywords = "Prompt Studio, Prompt 评测, 团队协作, 使用说明"

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_paragraph.add_run("Prompt Studio  |  团队使用说明")
    set_run_font(header_run, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_paragraph.add_run("内部协作资料  ·  ")
    set_run_font(footer_run, size=8.5, color=MUTED)
    add_page_number(footer_paragraph)

    cover_spacing = document.add_paragraph()
    cover_spacing.paragraph_format.space_after = Pt(68)
    cover_logo = document.add_paragraph()
    cover_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_logo.paragraph_format.space_after = Pt(18)
    cover_shape = cover_logo.add_run().add_picture(str(logo_path), width=Inches(1.12))
    cover_shape._inline.docPr.set("descr", "Prompt Studio 品牌 Logo")
    kicker = document.add_paragraph(style="Kicker")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.add_run("TEAM ENABLEMENT GUIDE")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("Prompt Studio")
    set_run_font(title_run, size=31, color=INK, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle_run = subtitle.add_run("团队协作 Prompt 评测与调试平台\n使用说明")
    set_run_font(subtitle_run, size=16, color=TEAL, bold=True)
    cover_intro = document.add_paragraph()
    cover_intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_intro.paragraph_format.space_after = Pt(30)
    cover_intro.paragraph_format.line_spacing = 1.35
    cover_intro_run = cover_intro.add_run("让 Prompt 的修改、对话效果、模型选择和团队意见\n都沉淀为可复用、可比较、可追溯的工作流")
    set_run_font(cover_intro_run, size=11, color=MUTED)
    metadata = document.add_table(rows=2, cols=3)
    set_table_geometry(metadata, [3120, 3120, 3120], indent=120)
    repeat_table_header(metadata.rows[0])
    metadata_values = [("适用对象", "产品 / 算法 / 内容 / 评测团队"), ("核心入口", "http://localhost:5173/"), ("文档版本", "2026.08")]
    for index, (label, value) in enumerate(metadata_values):
        header_cell = metadata.cell(0, index)
        value_cell = metadata.cell(1, index)
        accent = TEAL if index != 1 else CORAL
        set_cell_shading(header_cell, accent)
        set_cell_shading(value_cell, TEAL_SOFT if index != 1 else CORAL_SOFT)
        set_cell_border(header_cell, color=accent, size="6")
        set_cell_border(value_cell, color=accent, size="6")
        header_paragraph = header_cell.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_run = header_paragraph.add_run(label)
        set_run_font(header_run, size=8.5, color=WHITE, bold=True)
        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_run = value_paragraph.add_run(value)
        set_run_font(value_run, size=9, color=INK, bold=True)
    set_table_geometry(metadata, [3120, 3120, 3120], indent=120)
    document.add_page_break()

    add_heading(document, "0. 一页快速上手", 1)
    add_body(document, "Prompt Studio 是团队共同维护 Prompt、测试条件和评测证据的工作台。第一次使用不必理解全部模块，只要先完成下面五步。")
    add_callout(document, "五步完成一次评测", "调试 Prompt → 保存新版本 → 选择旧版和新版 → 运行固定测试集 → 逐条人工复核。先走通这条路径，再按需要学习模型卡、记忆和团队管理。")
    add_table(document, ["第几步", "从哪里进入", "只做这件事"], [
        ["1", "Prompt 调试器", "左右对话，找到比当前版本更好的 Prompt。"],
        ["2", "Prompt 版本", "把满意内容保存为新版本，不覆盖旧版本。"],
        ["3", "共享资源", "确认角色、世界观、测试集和模型卡已经准备好。"],
        ["4", "评测与结果", "选择旧版 / 新版，运行同一批测试输入。"],
        ["5", "版本评审", "逐条比较两侧回复，填写分数、证据和结论。"],
    ], [1100, 2500, 5760])
    add_heading(document, "按你的身份，只看需要的路径", 2)
    add_table(document, ["你是谁", "第一入口", "推荐阅读"], [
        ["负责人", "团队与权限", "第 1、4、5、7、10 节"],
        ["Prompt 编辑者", "Prompt 调试器", "第 2、3、4、5 节"],
        ["评测发起人", "评测与结果", "第 3、4、5、6 节"],
        ["人工评审人", "版本评审", "第 6、7 节"],
        ["体验参与者", "多人聊天房", "第 9 节"],
        ["只读观察者", "项目概览", "第 5、6、7 节"],
    ], [2100, 2800, 4460], accent=CORAL)
    add_heading(document, "先认识三个词", 2)
    add_bullet(document, "Prompt 版本：一次保存下来的 Prompt 正文。新改动要新建版本，旧版本不被覆盖。")
    add_bullet(document, "测试集：提前准备好的一批玩家输入和预期行为，用来反复检验新旧版本。")
    add_bullet(document, "评测证据：某条测试输入下，旧版回复、新版回复、自动评分和人工意见的完整记录。")
    add_callout(document, "页面里的英文", "Baseline 就是旧版回复，Candidate 就是新版回复；“冻结条件”表示这次评测中的角色、模型和测试输入不会中途变化。", BLUE_SOFT, BLUE)
    add_image(document, flow_path, 6.35, "图 1  Prompt Studio 从探索到发布的标准工作流")

    add_heading(document, "平台能力总览", 2)
    add_table(document, ["能力", "你可以做什么", "最终得到什么"], [
        ["Prompt 调试", "选模型、改 Prompt、调参数，左右窗口用同一输入对话。", "可比较的回复"],
        ["Prompt 版本", "从已有版本继续编辑，保存新版本，发布或切回旧版本。", "清晰的版本链"],
        ["共享资源", "维护角色、世界观、测试集、评分卡、模型卡、调度和记忆。", "统一测试条件"],
        ["旧版 / 新版评测", "只改变 Prompt，其他条件保持一致。", "逐条运行证据"],
        ["人工评审", "逐条看两侧回复，打分并写原因。", "团队结论"],
        ["模型横评", "固定 Prompt 和测试集，比较多个模型。", "模型选择依据"],
        ["聊天与记忆", "体验多人发言、@ 唤起和有无记忆差异。", "真实体验反馈"],
    ], [1900, 4500, 2960])

    add_heading(document, "一、进入平台与账号", 1)
    add_heading(document, "首次进入", 2)
    add_numbered(document, "打开平台首页，首次使用时会出现“初始化团队工作区”。创建首位负责人账号，填写姓名、账号和密码。")
    add_numbered(document, "账号仅支持 3 至 40 位字母或数字；密码至少 3 个字符。账号创建后，负责人可以在“团队与权限”中继续创建内部成员账号。")
    add_numbered(document, "登录成功后，浏览器会自动恢复当前会话。退出登录后，需要重新输入账号和密码。")
    add_callout(document, "安全提醒", "不要把模型 API Key 写进 Prompt、测试集或截图。模型卡支持服务端加密配置，团队成员只应使用平台里的模型卡，不要在聊天窗口粘贴密钥。", AMBER_SOFT, AMBER)
    add_heading(document, "日常进入", 2)
    add_body(document, "日常工作从“项目工作区”开始。页面左侧导航是项目内的功能入口，顶部显示当前项目、当前成员和刷新 / 退出操作。你只能看到自己被授权的项目和操作。")

    add_heading(document, "二、Prompt 调试器：先试，再保存", 1)
    add_body(document, "Prompt 调试器适合快速探索。左右两侧是两个独立实验条件，可以分别选择模型卡、编辑 Prompt 和调整采样参数；发送同一条用户输入后，平台会同时展示两侧回复。")
    add_numbered(document, "进入左侧“Prompt 调试器”，确认左右两侧的模型卡。模型卡包含供应商、协议、Base URL、模型名和可用参数。")
    add_numbered(document, "在 System Prompt 编辑区修改文本。变量会在发送前使用共享角色、世界观和对话历史展开；你看到的是最终请求对应的对话效果。")
    add_numbered(document, "在参数区调整 Temperature、Max Token 和供应商支持的其他参数。Claude 等模型不支持的参数不要强行填写，适配层会按协议处理。")
    add_numbered(document, "输入同一条消息发送。左右窗口各自保留回复历史，便于观察风格、指令遵循和上下文承接差异。")
    add_numbered(document, "确定方向后，不要直接覆盖旧版本。进入“Prompt 版本”，从当前内容新建一版，写清楚改动摘要。")
    add_callout(document, "判断标准", "调试器用于发现方向，不等于正式结论。正式结论要进入固定测试集，运行旧版 / 新版评测，并由团队逐案例复核。", TEAL_SOFT, TEAL)

    add_heading(document, "三、Prompt 版本：每次修改都留下版本", 1)
    add_body(document, "Prompt 版本是团队协作的基线。版本正文保存后成为不可变快照，后续修改应继续创建新版本，这样评测结果和人工意见不会被新内容污染。")
    add_numbered(document, "进入“Prompt 版本”，选中要维护的 Prompt 资产。先查看当前发布版本和版本时间线。")
    add_numbered(document, "点击“新建 Prompt 版本”，选择从哪个版本继续，并填写版本标题、变更说明和 Prompt 正文。")
    add_numbered(document, "保存后，回到版本详情检查正文、变量和共享资源引用。确认无误后再发布。")
    add_numbered(document, "需要回到旧版本时，在对应版本上执行发布操作。旧版本不会被删除，历史评测仍保留原始快照。")
    add_table(document, ["操作", "建议", "不要这样做"], [
        ["命名", "用 v2 · 减少追问、限制 NPC 抢话等可识别标题。", "只写“新版”“测试版”，让别人无法理解改了什么。"],
        ["变更说明", "说明改动目标、影响维度和预期风险。", "只复制 Prompt 正文，不写变更意图。"],
        ["发布", "先完成受控评测和人工复核，再发布为当前版本。", "调试器里感觉不错就直接发布。"],
        ["回滚", "发布旧版本，保留新版本继续分析。", "删除旧版本，破坏对比基线。"],
    ], [1900, 3700, 3760], accent=CORAL)

    add_heading(document, "四、共享资源：把评测条件集中管理", 1)
    add_body(document, "共享资源是两侧实验共同使用的基础条件。评测时，平台会把选中的角色、世界观、测试集、评分卡、模型、调度和记忆策略冻结到运行快照中。")
    add_table(document, ["资源", "配置重点", "什么时候修改"], [
        ["角色", "角色名、性格、人设边界；多人群聊要维护完整角色名册。", "角色设定发生业务变化时"],
        ["世界观", "故事背景、世界规则、地点、时间和不能违背的事实。", "剧情设定或产品世界观变化时"],
        ["测试集", "每条案例的玩家输入、上下文、期望行为和标签。", "补充典型场景、Bad Case 或回归案例时"],
        ["评分卡", "指令遵循、玩家行动权、剧情合理性、有趣性、人设、语气、世界观等维度。", "团队评分口径变化时"],
        ["模型卡", "协议、供应商、Base URL、模型名、参数和密钥引用。", "新增供应商、模型或接口配置时"],
        ["调度 / 记忆", "随机角色数量、@ 唤起、主动发言、记忆策略和历史窗口。", "多人机制或记忆方案变化时"],
    ], [1750, 4750, 2860])
    add_callout(document, "冻结原则", "一次正式评测开始后，不要中途修改共享资源。需要修改时，创建新资源版本或重新发起评测，让新旧运行保持可解释。", AMBER_SOFT, AMBER)

    add_heading(document, "五、正式评测：新旧 Prompt 如何对比", 1)
    add_body(document, "正式评测的基本单位是“一个 Prompt 版本 + 一组固定条件 + 一条测试集案例产生的对话证据”。平台只比较旧版和新版，不把 Builder A / B / C 暴露给协作者。")
    add_numbered(document, "进入“评测与结果”，选择同一个 Prompt 下的旧版和新版。")
    add_numbered(document, "选择测试集，确认测试集案例数量、角色、世界观、模型、历史窗口、调度、记忆和评分卡。")
    add_numbered(document, "运行前查看调用量估算和条件检查。缺少模型卡、API Key 或测试案例时，先回到共享资源修复。")
    add_numbered(document, "运行评测。每条案例会保存两侧原始回复、请求、响应、自动评分和失败原因；服务失败不会伪装成 0 分。")
    add_numbered(document, "先看总分和各维度差异，再打开逐案例结果。重点检查高分和低分案例是否符合实际体验。")
    add_numbered(document, "从完成的评测发起“版本评审”，把逐案例证据分配给评审者。")
    add_image(document, flow_path, 6.35, "图 2  正式评测应从共享条件冻结开始，并以逐案例证据结束")
    add_callout(document, "重要", "平均分只能告诉你变化方向，不能替代人工判断。一个严重的玩家行动权问题，可能被很多普通回复的平均分掩盖。", RED_SOFT, RED)

    add_heading(document, "六、版本评审：按一条一条的证据复核", 1)
    add_body(document, "版本评审不是给一个 Prompt 版本笼统打分，而是把 AI 生成出的每一轮对话证据发给评审者，要求评审者逐条对比 Baseline 与 Candidate，记录分数和理由。")
    add_numbered(document, "进入“版本评审”，选择一条已完成的旧版 / 新版评测，发起评审任务。")
    add_numbered(document, "指定项目内的评审者，并写明本轮评审关注的维度，例如玩家指令遵循、世界观遵循或群像互动。")
    add_numbered(document, "评审者打开案例，先阅读玩家输入和上下文，再左右对比两侧完整回复；不要只看 Candidate 的单侧表现。")
    add_numbered(document, "按评分卡逐维度打分，并写下可复核证据：哪句话遵循了要求，哪句话抢走了玩家决定权，哪处违反了世界观。")
    add_numbered(document, "提交通过、要求修改或仅评论的结论。后续修订会追加新记录，不会覆盖历史意见。")
    add_table(document, ["复核时看什么", "高质量证据示例", "低质量写法"], [
        ["玩家行动权", "玩家说“我先观察”，AI 只补充现场变化，没有替玩家决定结果。", "感觉没有抢戏。"],
        ["世界观遵循", "回复使用了测试集给出的地点和规则，没有新增未定义能力。", "世界观还可以。"],
        ["人设与语气", "莉亞保持温柔克制；诺亚使用计划和时间信息，语气可区分。", "角色很像本人。"],
        ["剧情合理性", "只推进一个线索，并且承接了上一轮已确认的事实。", "剧情比较顺。"],
    ], [2100, 4200, 3060], accent=BLUE)

    document.add_page_break()
    add_heading(document, "七、团队角色与推荐路径", 1)
    add_body(document, "平台提供四种项目角色。组织成员和项目成员是两层关系：先创建并激活账号，再把成员加入具体项目。权限由服务端校验，不能靠前端按钮隐藏来代替权限控制。")
    add_image(document, roles_path, 6.35, "图 3  四类团队角色的推荐使用路径")
    add_table(document, ["角色", "可以做什么", "推荐工作路径"], [
        ["负责人", "管理成员、项目、全部资产、发布和权限。", "创建项目 → 创建成员账号 → 分配项目角色 → 审批发布结论"],
        ["编辑者", "编辑 Prompt、共享资源、测试集，创建版本并运行评测。", "调试器 → 共享资源 → 新建版本 → 运行评测 → 发起评审"],
        ["评审者", "查看评测证据，逐案例评分、写证据、提交评审意见。", "版本评审 → 打开案例 → 对比回复 → 打分 → 提交结论"],
        ["观察者", "只读查看项目、版本和评测证据，可参与多人聊天体验。", "项目概览 → Prompt 版本 → 评测结果 → 多人聊天房"],
    ], [1500, 3800, 4060])

    add_heading(document, "八、模型卡与供应商接入", 1)
    add_body(document, "模型卡是团队共享的模型连接配置。它把协议差异收进供应商适配层，调试器和评测流程只需要选择模型卡，不需要每次重新拼接请求。")
    add_table(document, ["供应商 / 协议", "适用场景", "配置注意"], [
        ["OpenAI 兼容", "供应商提供 /v1/chat/completions 接口。", "配置协议、Base URL、模型名和服务端密钥。"],
        ["Cloudsway", "Cloudsway 专用 Chat Completions 请求。", "按 Cloudsway 适配器配置模型和 AK，不要手动改请求体。"],
        ["OpenRouter", "统一访问多个模型。", "可配置 X-Title 等扩展请求头；模型参数按被测模型支持范围填写。"],
        ["Anthropic Messages", "Claude 系列消息协议。", "不要发送 Claude 不支持的参数；以模型卡和适配器支持项为准。"],
    ], [2100, 3400, 3860], accent=AMBER)
    add_callout(document, "密钥管理", "模型卡可以引用服务端环境变量，也可以使用服务端加密保存。密钥原文不会返回浏览器，不会写入评测快照和审计记录。", TEAL_SOFT, TEAL)

    add_heading(document, "九、多人聊天房与记忆实验室", 1)
    add_heading(document, "多人聊天房", 2)
    add_numbered(document, "从工作区左侧进入“多人聊天房”，或打开加入页面输入昵称。进入后会看到左右 Prompt 对比窗口。")
    add_numbered(document, "在“角色配置”中维护角色和世界观，在 Builder 配置中维护两侧 Prompt；点击“同步配置”后再发送测试消息。")
    add_numbered(document, "普通消息、@ 唤起和 AI 主动发言按房间机制运行；@ 角色时，只让被点名角色进入回复线。")
    add_numbered(document, "需要重新开始时使用“清空记录”。该操作用于清空当前房间上下文并重启会话，不只是把页面上的消息隐藏。")
    add_heading(document, "记忆系统评测", 2)
    add_numbered(document, "进入“记忆实验室”，选择对照 Prompt、记忆更新频率、历史消息条数和 RAG Top K。")
    add_numbered(document, "同一条玩家输入会发送到“无记忆基线”和“记忆系统”两侧，观察角色状态、事件链、关键物品和检索命中。")
    add_numbered(document, "用典型的跨轮事实、人物关系、地点路线和玩家偏好测试记忆是否被正确保留，同时观察是否引入虚构信息。")

    add_heading(document, "十、团队协作约定与常见问题", 1)
    add_heading(document, "推荐约定", 2)
    for bullet in [
        "Prompt 改动必须新建版本，不直接覆盖已经发布的正文。",
        "每次评测尽量只改变一个主要变量，角色、世界观、测试集和模型保持一致。",
        "测试集要同时包含正常案例、边界案例和历史 Bad Case，避免只挑容易得分的输入。",
        "人工评审必须写可复核证据，优先引用具体回复和具体规则，不写“感觉好”“感觉差”。",
        "模型卡由负责人或编辑者维护，API Key 通过安全渠道配置，不在团队聊天中传播。",
        "正式发布前至少完成一次自动评测和一次人工逐案例复核。",
    ]:
        add_bullet(document, bullet)
    add_heading(document, "常见问题", 2)
    add_table(document, ["现象", "处理方式"], [
        ["平台显示不可用", "先确认前端 5173 和 API 3001 服务已启动，再点击重试；若为登录问题，重新登录。"],
        ["没有模型可选", "进入共享资源 → 模型卡，检查供应商、模型名、协议和密钥引用是否完整。"],
        ["评测很快结束", "检查测试集实际案例数量和运行状态；评测结果应能展开看到逐案例原始回复。"],
        ["AI 回复失败", "查看案例的失败原因和原始请求，不要把失败案例当成 0 分；先修模型卡再重新运行。"],
        ["成员看不到项目", "负责人需要先确认账号已激活，再在团队与权限中把成员加入当前项目。"],
        ["想比较两条 Prompt", "使用 Prompt 调试器做快速探索；要形成团队结论，必须保存版本并走旧版 / 新版评测。"],
    ], [2200, 7160], accent=TEAL)
    add_callout(document, "团队共识", "Prompt Studio 的价值不在于替团队做决定，而在于让每个决定都有清晰输入、可比较回复、可解释评分和可追溯版本。", CORAL_SOFT, CORAL)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
